# -*- coding: utf-8 -*-
"""Tests for the PPT deck evaluation module."""
import pytest

pytest.importorskip("pptx")

from fastapi.testclient import TestClient

import core.storage as storage
import core.ppt as ppt_core
import server as server_mod
from pptx import Presentation
from pptx.util import Inches, Pt


@pytest.fixture()
def deck_factory(tmp_path):
    def _make(name="deck.pptx", slides=3):
        prs = Presentation()
        layout = prs.slide_layouts[1]
        for i in range(slides):
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = f"第{i + 1}节标题"
            body = slide.placeholders[1]
            body.text = f"这是第{i + 1}页的正文内容，用于测试。" * 5
            for para in body.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(14)
                    run.font.name = "Calibri"
        path = tmp_path / name
        prs.save(str(path))
        return str(path)
    return _make


def _upload(client, deck_factory, name="deck.pptx", tag=""):
    path = deck_factory(name=name)
    with open(path, "rb") as f:
        r = client.post("/api/deck/upload", files={
            "file": (name, f.read(),
                     "application/vnd.openxmlformats-officedocument.presentationml.presentation")
        }, data={"tool_tag": tag})
    assert r.status_code == 200, r.text
    return r.json()


def test_deck_upload_dedupe_list_delete(isolated_db, deck_factory):
    c = TestClient(server_mod.server)
    up1 = _upload(c, deck_factory, tag="kimi")
    assert up1["duplicate"] is False and up1["slide_count"] == 3
    up2 = _upload(c, deck_factory, tag="kimi")
    assert up2["duplicate"] is True and up2["deck_id"] == up1["deck_id"]
    listed = c.get("/api/decks").json()
    assert len(listed) == 1 and listed[0]["tool_tag"] == "kimi"
    assert c.delete(f"/api/deck/{up1['deck_id']}").status_code == 200
    assert c.get("/api/decks").json() == []


def test_deck_upload_rejects_non_pptx(isolated_db, tmp_path):
    c = TestClient(server_mod.server)
    fake = tmp_path / "not.pptx"
    fake.write_bytes(b"stub")
    r = c.post("/api/deck/upload", files={"file": ("not.pptx", b"stub", "application/octet-stream")})
    assert r.status_code == 400


def test_deck_upload_rejects_corrupt_pptx(isolated_db):
    c = TestClient(server_mod.server)
    r = c.post("/api/deck/upload", files={
        "file": ("broken.pptx", b"PK\x03\x04garbage", "application/octet-stream")})
    assert r.status_code == 400


def test_p10_local_metric_persists(isolated_db, deck_factory):
    c = TestClient(server_mod.server)
    deck = _upload(c, deck_factory)
    r = c.post("/api/deck/evaluate/P10", data={
        "deck_id": deck["deck_id"], "base_url": "", "api_key": "", "model": ""})
    assert r.status_code == 200, r.text
    body = r.json()
    assert 0 <= body["value"] <= 10 and body["method"] == "objective"
    board = c.get("/api/deck/dashboard").json()
    assert len(board["decks"]) == 1
    assert board["decks"][0]["mean_scores"]["P10"] == body["value"]
    assert board["models"][0]["tool_tag"] == "未标注"


def test_evaluate_all_without_lmm_and_renderer(isolated_db, deck_factory, monkeypatch):
    monkeypatch.setattr(ppt_core, "detect_renderer",
                        lambda: {"backend": None, "detail": "none"})
    c = TestClient(server_mod.server)
    deck = _upload(c, deck_factory)
    r = c.post("/api/deck/evaluate-all", data={"deck_id": deck["deck_id"]})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["P10"]["value"] is not None
    for dim in ("P01", "P04", "P08", "P09"):
        assert body[dim]["value"] is None and "视觉模型" in body[dim]["note"]
    for dim in ("P05", "P06", "P07"):
        assert body[dim]["value"] is None and body[dim]["note"] == "NO_RENDERER"
    # only successful dims persist
    board = c.get("/api/deck/dashboard").json()
    assert set(board["decks"][0]["mean_scores"]) == {"P10"}


def test_deck_design_dim_502_when_no_renderer(isolated_db, deck_factory, monkeypatch):
    monkeypatch.setattr(ppt_core, "detect_renderer",
                        lambda: {"backend": None, "detail": "none"})
    c = TestClient(server_mod.server)
    deck = _upload(c, deck_factory)
    r = c.post("/api/deck/evaluate/P05", data={"deck_id": deck["deck_id"]})
    assert r.status_code == 502 and "NO_RENDERER" in r.json()["detail"]


def test_deck_slide_404_and_traversal_guard(isolated_db, deck_factory):
    c = TestClient(server_mod.server)
    deck = _upload(c, deck_factory)
    assert c.get(f"/api/deck/{deck['deck_id']}/slide/1.png").status_code == 404
    assert c.get("/api/deck/..%2F..%2Fevil/slide/1.png").status_code == 404
    assert c.get("/api/deck/zz/slide/1.png").status_code == 404


def test_get_deck_path_rejects_traversal(isolated_db, deck_factory):
    deck = _upload(TestClient(server_mod.server), deck_factory)
    assert storage.get_deck_path(deck["deck_id"])
    assert storage.get_deck_path("../escape") is None
    assert storage.get_deck_path("no_such_id") is None
